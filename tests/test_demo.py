from pathlib import Path
import unittest, tempfile, sys, csv
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from scripts.crear_demo import crear


class DemoTest(unittest.TestCase):
    def test_fuentes_y_wbs(self):
        with tempfile.TemporaryDirectory() as tmp:
            root=Path(tmp);path=crear(root)
            with path.open(encoding='utf-8-sig',newline='') as f:
                reader=csv.DictReader(f,delimiter=';');rows=list(reader)
                self.assertEqual(len(reader.fieldnames),49)
            self.assertEqual(len(rows),6)
            codigos={r['WBS'] for r in rows};total={w:0 for w in codigos}
            files=list((root/'avance_real').glob('*.xlsx'));self.assertEqual(len(files),4)
            for file in files:
                wb=load_workbook(file,data_only=True)
                for r in wb['DATOS'].iter_rows(min_row=5,values_only=True):
                    self.assertIn(r[2],codigos);total[r[2]]+=r[7]
                wb.close()
            self.assertTrue(all(v==70 for v in total.values()))
            with self.assertRaises(SystemExit):crear(root)
